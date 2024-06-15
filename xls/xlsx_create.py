from random import randint
import openpyxl
from openpyxl.chart import BarChart, Reference, PieChart


file = openpyxl.Workbook()
sheet = file.active

for i in range(1, 11):
    sheet.cell(row=i, column=1).value = randint(1, 11)

chart = Reference(sheet, min_col=1, min_row=1, max_row=10)

bar_chart = BarChart()

bar_chart.add_data(chart, titles_from_data=False)

sheet.add_chart(bar_chart, 'D5')

file.save('sample.xlsx')
